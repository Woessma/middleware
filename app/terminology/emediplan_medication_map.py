import os
from functools import lru_cache

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover
    load_workbook = None


DEFAULT_PHARMACODE_XLSX = os.getenv(
    "EMEDIPLAN_PHARMACODE_XLSX",
    "/app/terminology/data/20260901__MappingPharmacodeToGtinToSwissmedicDescription.xlsx",
)

DEFAULT_PRODUCTNUMBER_XLSX = os.getenv(
    "EMEDIPLAN_PRODUCTNUMBER_XLSX",
    "/app/terminology/data/20260701__MappingProductNumberToSwissmedicDescription.xlsx",
)

COMPENDIUM_SEARCH_URL = "https://compendium.ch/de/search"


def _clean(value):
    if value is None:
        return None

    text = str(value).strip()
    return text or None


def _normalize_header(header):
    if header is None:
        return ""

    return (
        str(header)
        .strip()
        .lower()
        .replace(" ", "")
        .replace("_", "")
        .replace("-", "")
    )


def _select_column(header_index, *candidates):
    for candidate in candidates:
        normalized = _normalize_header(candidate)

        if normalized in header_index:
            return header_index[normalized]

    return None


def _iter_rows_from_xlsx(path):
    if not path or not os.path.exists(path):
        return []

    if load_workbook is None:
        return []

    workbook = load_workbook(path, data_only=True, read_only=True)

    try:
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if not rows:
        return []

    headers = rows[0]
    header_index = {
        _normalize_header(name): idx
        for idx, name in enumerate(headers)
        if _normalize_header(name)
    }

    return rows[1:], header_index


@lru_cache(maxsize=1)
def _load_mapping():
    mapping = {
        "pharmacode": {},
        "product_number": {},
        "gtin": {},
    }

    parsed = _iter_rows_from_xlsx(DEFAULT_PHARMACODE_XLSX)

    if parsed:
        rows, header_index = parsed

        pharmacode_col = _select_column(header_index, "Pharmacode")
        gtin_col = _select_column(header_index, "GTIN")
        product_number_col = _select_column(
            header_index,
            "Swissmedic-Nr + Packungscode",
            "SwissmedicNr+Packungscode",
            "ProductNumber",
        )
        description_col = _select_column(
            header_index,
            "Swissmedic Bezeichnung des Arzneimittels",
            "SwissmedicDescription",
            "Description",
        )
        form_col = _select_column(
            header_index,
            "Form",
            "DosageForm",
            "GalenicForm",
            "Darreichungsform",
        )
        ingredient_col = _select_column(
            header_index,
            "Ingredient",
            "ActiveIngredient",
            "Wirkstoff",
            "Substance",
        )
        strength_col = _select_column(
            header_index,
            "Strength",
            "Stength",
            "Starke",
            "Stärke",
            "DosageStrength",
        )
        atc_col = _select_column(header_index, "ATC")
        dispensing_category_col = _select_column(
            header_index,
            "Abgabekategorie",
            "DispensingCategory",
        )

        for row in rows:
            pharmacode = _clean(row[pharmacode_col]) if pharmacode_col is not None else None
            gtin = _clean(row[gtin_col]) if gtin_col is not None else None
            product_number = (
                _clean(row[product_number_col]) if product_number_col is not None else None
            )
            description = (
                _clean(row[description_col]) if description_col is not None else None
            )
            form = _clean(row[form_col]) if form_col is not None else None
            ingredient = _clean(row[ingredient_col]) if ingredient_col is not None else None
            strength = _clean(row[strength_col]) if strength_col is not None else None
            atc = _clean(row[atc_col]) if atc_col is not None else None
            dispensing_category = (
                _clean(row[dispensing_category_col])
                if dispensing_category_col is not None
                else None
            )

            row_data = {
                "gtin": gtin,
                "product_number": product_number,
                "description": description,
                "form": form,
                "ingredient": ingredient,
                "strength": strength,
                "atc": atc,
                "dispensing_category": dispensing_category,
            }

            if pharmacode:
                mapping["pharmacode"][pharmacode] = row_data

            if gtin:
                mapping["gtin"][gtin] = row_data

            if product_number:
                mapping["product_number"].setdefault(product_number, row_data)

    parsed = _iter_rows_from_xlsx(DEFAULT_PRODUCTNUMBER_XLSX)

    if parsed:
        rows, header_index = parsed

        product_number_col = _select_column(
            header_index,
            "Swissmedic-Nr + Packungscode",
            "SwissmedicNr+Packungscode",
            "ProductNumber",
        )
        description_col = _select_column(
            header_index,
            "Swissmedic Bezeichnung des Arzneimittels",
            "SwissmedicDescription",
            "Description",
        )
        gtin_col = _select_column(header_index, "GTIN")
        form_col = _select_column(
            header_index,
            "Form",
            "DosageForm",
            "GalenicForm",
            "Darreichungsform",
        )
        ingredient_col = _select_column(
            header_index,
            "Ingredient",
            "ActiveIngredient",
            "Wirkstoff",
            "Substance",
        )
        strength_col = _select_column(
            header_index,
            "Strength",
            "Stength",
            "Starke",
            "Stärke",
            "DosageStrength",
        )
        atc_col = _select_column(header_index, "ATC")
        dispensing_category_col = _select_column(
            header_index,
            "Abgabekategorie",
            "DispensingCategory",
        )

        for row in rows:
            product_number = (
                _clean(row[product_number_col]) if product_number_col is not None else None
            )
            description = (
                _clean(row[description_col]) if description_col is not None else None
            )
            gtin = _clean(row[gtin_col]) if gtin_col is not None else None
            form = _clean(row[form_col]) if form_col is not None else None
            ingredient = _clean(row[ingredient_col]) if ingredient_col is not None else None
            strength = _clean(row[strength_col]) if strength_col is not None else None
            atc = _clean(row[atc_col]) if atc_col is not None else None
            dispensing_category = (
                _clean(row[dispensing_category_col])
                if dispensing_category_col is not None
                else None
            )

            row_data = {
                "gtin": gtin,
                "product_number": product_number,
                "description": description,
                "form": form,
                "ingredient": ingredient,
                "strength": strength,
                "atc": atc,
                "dispensing_category": dispensing_category,
            }

            if product_number:
                mapping["product_number"][product_number] = row_data

            if gtin:
                mapping["gtin"].setdefault(gtin, row_data)

    return mapping


def resolve_medication_identity(id_type, raw_id):
    med_id = _clean(raw_id)

    if not med_id:
        return {
            "id": None,
            "id_type": id_type,
            "display": None,
            "gtin": None,
            "product_number": None,
            "form": None,
            "ingredient": None,
            "strength": None,
            "atc": None,
            "dispensing_category": None,
            "source": None,
        }

    mapping = _load_mapping()

    if id_type == 3:
        found = mapping["pharmacode"].get(med_id) or {}
    elif id_type == 4:
        found = mapping["product_number"].get(med_id) or {}
    elif id_type == 2:
        found = mapping["gtin"].get(med_id) or {}
    else:
        found = {}

    resolved_id = med_id
    resolved_type = id_type

    # Prefer canonical GTIN if known.
    gtin = _clean(found.get("gtin"))
    if gtin:
        resolved_id = gtin
        resolved_type = 2

    source_reference = f"{COMPENDIUM_SEARCH_URL}?q={med_id}&type=Default"

    return {
        "id": resolved_id,
        "id_type": resolved_type,
        "display": _clean(found.get("description")),
        "gtin": gtin,
        "product_number": _clean(found.get("product_number")),
        "form": _clean(found.get("form")),
        "ingredient": _clean(found.get("ingredient")),
        "strength": _clean(found.get("strength")),
        "atc": _clean(found.get("atc")),
        "dispensing_category": _clean(found.get("dispensing_category")),
        "source": {
            "system": "https://compendium.ch",
            "reference": source_reference,
            "display": "Compendium",
        },
    }
